"""Fixture resume and job posting for Phase 1 ladder rung 5 (docs/08 Q-1, F-3, F-5).

A fictional person (Jordan Avery Example) and a fictional employer (Northwind Data Labs).
Writes into <out-dir>:
  resume.docx       the resume (python-docx)
  resume.pdf        the same resume (reportlab)
  job-posting.md    the posting
and into <key-dir> (keep it outside any folder the session can see):
  answer-key.md     the planted mismatches, for scoring only

    python -I bench/fixtures/resume_fixture.py <out-dir> <key-dir>
"""

import os
import sys

NAME = "Jordan Avery Example"
CONTACT = "jordan.example@example.com | Springfield, USA"

SUMMARY = (
    "Data engineer with 3 years of experience building batch pipelines in Python and SQL "
    "on Google Cloud Platform."
)
EXPERIENCE = [
    ("Data Engineer, Contoso Analytics (2023 to present)", [
        "Built daily batch ETL in Python and Airflow loading 2 TB/day into BigQuery.",
        "Wrote dbt models and SQL tests for finance reporting.",
        "Packaged jobs in Docker and deployed them on Cloud Run.",
    ]),
    ("Junior Data Analyst, Fabrikam Retail (2022 to 2023)", [
        "Built Tableau dashboards for store operations.",
        "Automated weekly Excel reports with pandas.",
    ]),
]
SKILLS = "Python, SQL, Airflow, dbt, BigQuery, Docker, Google Cloud Platform, pandas, Tableau"
EDUCATION = "B.S. Statistics, Springfield State University, 2022"

POSTING = """# Senior Data Engineer, Northwind Data Labs (fictional)

Location: Remote (US)

## About the role
You will own our streaming data platform and mentor two junior engineers.

## Requirements
- 5+ years of professional data engineering experience
- Strong Python and SQL
- Streaming systems: Apache Kafka and Apache Flink or Spark Structured Streaming
- AWS (S3, Glue, Redshift); infrastructure as code with Terraform
- Kubernetes in production
- Experience mentoring or leading engineers

## Nice to have
- Airflow or Dagster
- dbt
- Go or Scala

## Education
- Degree in Computer Science or a related field, or equivalent experience
"""

ANSWER_KEY = """# Planted mismatches (scoring only)
- experience: posting 5+ years, resume about 3 years (2022 junior analyst, 2023 data engineer)
- streaming: posting Kafka + Flink/Spark streaming, resume batch only
- cloud: posting AWS (S3, Glue, Redshift), resume GCP (BigQuery, Cloud Run)
- IaC: posting Terraform, resume none
- Kubernetes: posting production Kubernetes, resume Docker + Cloud Run only
- mentoring/leading: posting requires it, resume shows none
- senior title: posting Senior, resume mid-level
- matches: Python, SQL, Airflow (nice to have), dbt (nice to have), degree in Statistics (related field)
- not present in resume: Go or Scala
"""


def write_docx(path: str) -> None:
    import docx

    d = docx.Document()
    d.add_heading(NAME, level=0)
    d.add_paragraph(CONTACT)
    d.add_heading("Summary", level=1)
    d.add_paragraph(SUMMARY)
    d.add_heading("Experience", level=1)
    for role, bullets in EXPERIENCE:
        d.add_heading(role, level=2)
        for b in bullets:
            d.add_paragraph(b, style="List Bullet")
    d.add_heading("Skills", level=1)
    d.add_paragraph(SKILLS)
    d.add_heading("Education", level=1)
    d.add_paragraph(EDUCATION)
    d.save(path)


def write_pdf(path: str) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate

    s = getSampleStyleSheet()
    flow = [Paragraph(NAME, s["Title"]), Paragraph(CONTACT, s["Normal"]),
            Paragraph("Summary", s["Heading1"]), Paragraph(SUMMARY, s["Normal"]),
            Paragraph("Experience", s["Heading1"])]
    for role, bullets in EXPERIENCE:
        flow.append(Paragraph(role, s["Heading2"]))
        flow.append(ListFlowable([ListItem(Paragraph(b, s["Normal"])) for b in bullets],
                                 bulletType="bullet"))
    flow += [Paragraph("Skills", s["Heading1"]), Paragraph(SKILLS, s["Normal"]),
             Paragraph("Education", s["Heading1"]), Paragraph(EDUCATION, s["Normal"])]
    SimpleDocTemplate(path, pagesize=letter).build(flow)


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    out, key_dir = sys.argv[1], sys.argv[2]
    os.makedirs(out, exist_ok=True)
    os.makedirs(key_dir, exist_ok=True)
    write_docx(os.path.join(out, "resume.docx"))
    write_pdf(os.path.join(out, "resume.pdf"))
    with open(os.path.join(out, "job-posting.md"), "w", encoding="utf-8") as f:
        f.write(POSTING)
    with open(os.path.join(key_dir, "answer-key.md"), "w", encoding="utf-8") as f:
        f.write(ANSWER_KEY)
    print(f"wrote resume.docx, resume.pdf, job-posting.md to {out}; answer-key.md to {key_dir}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
